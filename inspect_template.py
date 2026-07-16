import openpyxl
wb = openpyxl.load_workbook('SMMS_Master_Data.xlsx', data_only=True)
print('Sheets:', wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    print('\n===', name, '===')
    for i, row in enumerate(rows[:10]):
        print(i+1, row)